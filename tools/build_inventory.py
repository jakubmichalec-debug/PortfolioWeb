import datetime
import io
import os
import re
import sys
import zipfile
from collections import defaultdict

sys.path.insert(0, os.path.dirname(__file__))
from helpers import geometry_key, montage, read_stl, render_stl, render_svg  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from PIL import Image  # noqa: E402

SRC = "C:/Users/jakub/Desktop/KdG/fablab"
OUT = "C:/Users/jakub/Desktop/PORTFOLIO/pinball-parts-inventory.xlsx"
PP = "Prototyping Pinball"

# ---------------------------------------------------------------- groups
BOX, WALL, FLIP, TARG, LOAD, SLID, BRAK, UNCL = (
    "Box & playfield (wood)",
    "Walls & lanes (wood)",
    "Flippers & drive",
    "Targets",
    "Ball loading",
    "Slide, barrier & small parts",
    "Brackets",
    "Unclear - please check",
)

ROWS = []


def R(group, comp, ver, stems, desc, how):
    ROWS.append(dict(group=group, comp=comp, ver=ver, stems=stems, desc=desc, how=how))


# ---- Box & playfield
R(BOX, "Box (early)", "BoxProject (Feb)", ["BoxProject"], "First simple open box, made before the finger-joint version.", "Laser cut")
R(BOX, "Box finger-joint test", "Bottom finger (Feb)", ["bocx bottomfinger", "BottomBox Finger"], "Early test of a finger-joint bottom panel.", "Laser cut")
R(BOX, "Box (finger joints)", "boxProjectFinger", ["boxProjectFinger"], "Main wooden box with finger joints: back, bottom, front, left/right sides, long and short legs. 6 mm MDF, laser cut (report).", "Laser cut")
R(BOX, "Box side panel", "left side (SVG)", ["left side"], "Side panel outline. SVG only, no FreeCAD source found.", "Laser cut")
R(BOX, "Box side panel", "Dimensions image", ["SideBoxDimensions"], "Reference picture with the side dimensions.", "Reference")
R(BOX, "Playfield plate", "middle part", ["middle part"], "First outline of the playfield plate.", "Laser cut")
R(BOX, "Playfield plate", "Sketch base Middle", ["Sketch base Middle"], "Second sketch of the playfield plate.", "Laser cut")
R(BOX, "Playfield plate", "Sketch Middle New", ["Sketch Middle New"], "Third sketch of the playfield plate.", "Laser cut")
R(BOX, "Playfield plate", "SketchMiddleNew", ["SketchMiddleNew"], "Playfield layout with flippers, targets, walls and slide. Laser-engraved as an assembly guide (report).", "Laser cut")
R(BOX, "Playfield plate", "Unnamed1 exports", ["Unnamed1"], "Exports of body sketch, bottom cover, load-shaft cover and a triangle target STL. The FreeCAD file was renamed, source not found.", "Laser cut + 3D print")
R(BOX, "Load-shaft cover", "CoverShaftLoad", ["CoverShaftLoad"], "Cover over the ball-loading shaft.", "Laser cut")

# ---- Walls & lanes
R(WALL, "Curved wall (inner)", "innercurve", ["innercurve"], "Inner curved wall pieces (8 strips in the SVG).", "Laser cut")
R(WALL, "Curved wall (outer)", "outercurve", ["outercurve"], "Outer curved wall pieces.", "Laser cut")
R(WALL, "Loading wall", "WallLoading", ["WallLoading"], "Wall along the ball-loading lane.", "Laser cut")
R(WALL, "Top triangle wall", "WallTopTriangle", ["WallTopTriangle"], "Wall piece at the top triangle of the playfield.", "Laser cut")
R(WALL, "Left wall (arms)", "LeftWallArms", ["LeftWallArms"], "Wooden wall next to the flipper arms.", "Laser cut")
R(WALL, "Flipper side wood", "Left (LeftSideRucickyDrevo)", ["LeftSideRucickyDrevo"], "Wood side piece for the left flipper ('drevo' = wood).", "Laser cut")
R(WALL, "Flipper side wood", "Right (RightSideRucickyDrevo)", ["RightSideRucickyDrevo"], "Wood side piece for the right flipper.", "Laser cut")
R(WALL, "Wall left flipper", "WallLeftRucicky", ["WallLeftRucicky"], "Wall outline near the left flipper. The STL is empty, only the SVG has data.", "Laser cut")
R(WALL, "Loading lane wood", "LoadingWood", ["LoadingWood"], "First wooden guide wall for the ball storage lane (report).", "Laser cut")
R(WALL, "Loading lane wood", "LoadingWoodLong", ["LoadingWoodLong"], "Long guide wall of the ball lane.", "Laser cut")
R(WALL, "Loading lane wood", "LoadingWoodShort", ["LoadingWoodShort"], "Short guide wall of the ball lane.", "Laser cut")
R(WALL, "Ball chamber wood", "ChamberBottomShort", ["ChamberBottomShort"], "Short bottom piece of the ball chamber.", "Laser cut")
R(WALL, "Ball chamber wood", "ChamberLongWood", ["ChamberLongWood"], "Long wooden piece of the ball chamber.", "Laser cut")
R(WALL, "Ball chamber wood", "ChamberTopCurve", ["ChamberTopCurve"], "Curved top piece of the ball chamber.", "Laser cut")

# ---- Flippers & drive
R(FLIP, "Flipper lever", "RucickaBolt", ["RucickaBolt"], "First flipper: nut-and-bolt pivot (report). Two shapes exported: Rucicka and RucickaNew.", "3D print")
R(FLIP, "Flipper lever", "RucickaNovaV3 (STL only)", ["RucickaNovaV3"], "'Nova' = new flipper design. STL export only, no FreeCAD source.", "3D print")
R(FLIP, "Flipper lever", "RucickaBoltV4", ["RucickaBoltV4"], "Bolt-pivot flipper, 4th version.", "3D print")
R(FLIP, "Flipper lever", "RucickaNovaV4 (STL only)", ["RucickaNovaV4"], "'Nova' flipper, 4th version. STL export only.", "3D print")
R(FLIP, "Flipper lever", "RucickaNova", ["RucickaNova"], "Newest flipper file. Final flipper used a steel rod with a D-shaped end (report).", "3D print")
R(FLIP, "Flipper drive adapter", "V1 (BottomArmAdapter)", ["BottomArmAdapter"], "Adapter under the playfield linking the solenoid arm to the flipper rod (report).", "3D print")
R(FLIP, "Flipper drive adapter", "V2", ["BottomArmAdapterV2"], "Adapter under the playfield, 2nd version.", "3D print")
R(FLIP, "Flipper drive adapter", "V3", ["BottomArmAdapterV3"], "Adapter under the playfield, 3rd version.", "3D print")
R(FLIP, "Flipper drive adapter", "V4", ["BottomArmAdapterV4"], "Adapter under the playfield, 4th version (split into Adapter, Body and Connector).", "3D print")
R(FLIP, "Flipper drive adapter", "V5 (BottomAdapterV5)", ["BottomAdapterV5"], "Adapter under the playfield, 5th version. No STL exported.", "3D print")
R(FLIP, "Flipper solenoid assembly", "AdapterArmsV1", ["AdapterArmsV1"], "Base + holder, arm, sleeve, circle and connector for the flipper mechanism.", "3D print")
R(FLIP, "Flipper solenoid assembly", "AdapterRucickyL", ["AdapterRucickyL"], "Reworked copy of AdapterArmsV1 for the left flipper: same parts, but smaller in one direction (base 22 vs 27 mm, sleeve 25 vs 29 mm). This is the one in your MondayPrint batch.", "3D print")
R(FLIP, "Solenoid adapter (mini)", "AdapterSolenoidMini", ["AdapterSolenoidMini"], "Small adapter for a mini solenoid.", "3D print")
R(FLIP, "Solenoid bracket (big)", "BracketSolenoidBig", ["BracketSolenoidBig"], "Bracket for the big solenoid.", "3D print")
R(FLIP, "Arm-solenoid connector", "STL only", ["ArmSolenoidConnector"], "Connector between the arm and the solenoid (8 x 28 x 2.5 mm). STL export only.", "3D print")

# ---- Targets
R(TARG, "Hit target (top)", "V1", ["HitTargetTopV1"], "Target with a movable piece and metal contacts behind it; closing the contacts scores points (report).", "3D print")
R(TARG, "Hit target (top)", "V2", ["HitTargetTopV2"], "Top target, 2nd version.", "3D print")
R(TARG, "Hit target (top)", "V3", ["HitTargetTopV3"], "Top target, 3rd version (inside and outside target pieces).", "3D print")
R(TARG, "Hit target (top)", "V4", ["HitTargetTopV4"], "Top target, 4th version: adapter holder, inside/outside target and triangle hit target.", "3D print")
R(TARG, "Hit target (top)", "V5", ["HitTargetTopV5"], "Top target, 5th version. No STL exported.", "3D print")
R(TARG, "Hit target (mid)", "V1", ["HitTargetMidV1"], "Middle target with adapter and triangle.", "3D print")
R(TARG, "Hit target (mid)", "V2", ["HitTargetMidV2"], "Middle target, 2nd version, with stopper and triangle.", "3D print")
R(TARG, "Hit target (corner)", "V1", ["HitTargetCornerV1"], "Corner target: holder, inside/outside target and triangle.", "3D print")
R(TARG, "Hit target (slide)", "V1", ["HitTargetSlide"], "Target at the slide: holder, target, target insert and angle piece ('uholnik').", "3D print")
R(TARG, "Drop target", "V1", ["DropTarget"], "Drop target mounted through the playfield.", "3D print")
R(TARG, "Drop target", "V2", ["DropTargetV2"], "Drop target, 2nd version.", "3D print")

# ---- Ball loading
R(LOAD, "Loading system diagram", "LoadingAll.png", ["LoadingAll"], "Your overview picture: loading ramp, hook, solenoid, microswitch, ramp holder, manual plunger.", "Reference")
R(LOAD, "Loading (early parts)", "LoadingBottomLeg", ["LoadingBottomLeg"], "Early loading system part, 11 Apr.", "3D print")
R(LOAD, "Loading (early parts)", "LoadingTopLeg", ["LoadingTopLeg"], "Early loading system part, 11 Apr.", "3D print")
R(LOAD, "Loading (early parts)", "LoadingLock", ["LoadingLock"], "Early loading lock, 12 Apr.", "3D print")
R(LOAD, "Loading hook", "HookKetchV1", ["HookKetchV1"], "First hook that stops the balls ('ketch' = catch).", "3D print")
R(LOAD, "Loading hook", "HookKetchV2 (STL only)", ["HookKetchV2"], "Hook, 2nd version. STL export only.", "3D print")
R(LOAD, "Loading hook", "HookLoadingV1", ["HookLoadingV1"], "Hook for the loading system, 3 May.", "3D print")
R(LOAD, "Loading hook", "HookV3", ["HookV3"], "Curved hook driven by a solenoid; releases one ball at a time (report).", "3D print")
R(LOAD, "Ball shoot platform", "BallLoadingShootPlatform", ["BallLoadingShootPlatform"], "Platform where the ball is loaded and shot out.", "3D print")
R(LOAD, "RFID holder", "RFIDHolder", ["RFIDHolder"], "Holder for the RFID reader that starts a game and resets the score (report).", "3D print")

# ---- Slide, barrier, small parts
R(SLID, "Slide", "SlideV1", ["SlideV1"], "Ball slide, printed in two halves and glued (report): Body, Left, Right.", "3D print")
R(SLID, "Barrier ('Zabrana')", "ZabranaV1", ["ZabranaV1"], "Cylindrical barrier with mounting hole that guides the ball (report).", "3D print")
R(SLID, "Limiter ('Obmedzovac')", "Obmedzovac", ["Obmedzovac", "addaadad"], "Small cylinder, 'obmedzovac' = limiter/stopper. addaadad.stl (a 20 x 20 x 35 mm cylinder) is probably its export.", "3D print")
R(SLID, "Washer ('Podlozka')", "Podlozka", ["Podlozka", "Washer"], "Washer/spacer (38 mm across, 5 mm thick). Washer.stl is its export.", "3D print")

# ---- Brackets
R(BRAK, "L bracket (large)", "V1", ["LBracketsV1"], "Corner bracket that holds the wooden box together (report).", "3D print")
R(BRAK, "L bracket (small)", "V1", ["LbracketSmallV1"], "Second, smaller bracket design (report mentions two).", "3D print")

# ---- Unclear
R(UNCL, "Servo adapter", "DervoV1", ["DervoV1"], "Servo motor holder ('Dervo' is probably 'Servo'). Pinball uses solenoids, so maybe an early idea or a launcher leftover.", "3D print?")
R(UNCL, "Servo adapter", "ServoAdapterV2", ["ServoAdapterV2"], "Servo motor adapter, 2nd version. Same question as above.", "3D print?")
R(UNCL, "Launcher adapter (maybe)", "ADAPTER2 (sliced 28 Apr)", ["ADAPTER2"], "Two print files sliced 28 Apr. ADAPTER2 also exists in the launcher folder, so it may not be pinball.", "3D print")
R(UNCL, "Scratch / test file", "df", ["df"], "Unnamed small U-shaped clip (19 May).", "?")
R(UNCL, "Scratch / test file", "sfdf", ["sfdf"], "Unnamed sketch of playfield walls (19 May).", "?")
R(UNCL, "Scratch / test file", "Test", ["Test"], "Small flat clip (2 May).", "?")
R(UNCL, "Scratch / test file", "Test21", ["Test21"], "Angled bracket with two holes (2 May).", "?")
R(UNCL, "Scratch / test file", "Test30hole (STL)", ["Test30hole"], "STL only, 7 May. Maybe a hole-size test print.", "3D print?")
R(UNCL, "Scratch / test file", "Drawing2", ["Drawing2"], "Outline sketch of an adapter (5 Jun).", "?")
R(UNCL, "Scratch / test file", "Unnamed2 (Feb)", ["Unnamed2"], "Dimension sketch 500 x 125 mm from 28 Feb; a tiny STL was exported from it on 19 May.", "?")
R(UNCL, "Scratch / test file", "Unnamed3 (SVG, Feb)", ["Unnamed3"], "Rectangle outline SVG from 28 Feb.", "Laser cut?")

# files that exist in the folder but are clearly not the pinball (shown on the 'Not pinball' tab)
NON_PINBALL = [
    ("fablab.ino", "Arduino dice thrower with an LCD (Nov 2025)", "Small class exercise, no relation to the pinball."),
    ("bench01.FCStd + bench01-Pad004_tab.dxf", "Bench model + DXF (Feb 2026)", "FreeCAD tutorial-style project from before the pinball started."),
    ("basePartTurotial.FCStd", "Bracket from a tutorial (27 Feb)", "Tutorial file ('Turotial')."),
    ("RandomTurorial1.FCStd", "Cylinder on a base plate (1 May)", "Tutorial file."),
    ("PracticePart10.FCStd", "Practice part (5 Jun)", "Practice part, dated after the machine was built."),
    ("PracticePart11.FCStd", "Practice part (5 Jun)", "Practice part, dated after the machine was built."),
    ("adadad.FCStd", "Small wedge (27 Feb)", "Same day as the tutorial file, looks like a keyboard-mash test."),
    ("FabLAB All canon/ (folder)", "Ping pong launcher files (Dec 2025 - Jan 2026)", "The launcher project you decided to leave out of the portfolio."),
    ("Fusion challange/ (folder)", "Lost & Found locker + Mental Hospital Museum", "Different projects."),
    ("INenOUT/ (folder)", "Arduino class material and exercises", "Different course material."),
    ("Prototyping/ (folder)", "Paper cube lamp, Billboard scene, old course portfolio", "Different projects."),
]
NON_PINBALL_PATTERNS = [
    r"^fablab\.ino$", r"^bench01", r"^basePartTurotial", r"^RandomTurorial1", r"^PracticePart1[01]", r"^adadad\.",
]

# ---------------------------------------------------------------- scan files
ALL = []  # dicts: rel, base, sub, mtime, size
for f in sorted(os.listdir(SRC)):
    p = os.path.join(SRC, f)
    if os.path.isfile(p):
        ALL.append(dict(rel=f, base=f, sub="", path=p))
for root, _, files in os.walk(os.path.join(SRC, PP)):
    for f in files:
        p = os.path.join(root, f)
        rel = os.path.relpath(p, SRC).replace("\\", "/")
        ALL.append(dict(rel=rel, base=f, sub=os.path.relpath(root, os.path.join(SRC, PP)).replace("\\", "/").replace(".", ""), path=p))
ALL = [a for a in ALL if not a["base"].lower().endswith(".fcbak")]
for a in ALL:
    st = os.stat(a["path"])
    a["mtime"] = datetime.date.fromtimestamp(st.st_mtime)
    a["size"] = st.st_size
    a["ext"] = os.path.splitext(a["base"])[1].lower()

assigned = {}
for i, row in enumerate(ROWS):
    row["files"] = []
    for stem in row["stems"]:
        rx = re.compile("^" + re.escape(stem) + r"(?=[-._])")
        for a in ALL:
            if rx.match(a["base"]):
                if a["rel"] in assigned and assigned[a["rel"]] != i:
                    raise SystemExit(f"file {a['rel']} matched by two rows: {assigned[a['rel']]} and {i}")
                if a["rel"] not in assigned:
                    assigned[a["rel"]] = i
                    row["files"].append(a)

unassigned = [a for a in ALL if a["rel"] not in assigned]
leftover = [a["rel"] for a in unassigned if not any(re.search(p, a["base"]) for p in NON_PINBALL_PATTERNS)]
if leftover:
    raise SystemExit("Unassigned files:\n" + "\n".join(leftover))
empty_rows = [r["ver"] for r in ROWS if not r["files"]]
if empty_rows:
    raise SystemExit(f"rows without files: {empty_rows}")
print(f"{len(ROWS)} rows, {len(assigned)} files assigned, {len(unassigned)} deliberately left out")

# ---------------------------------------------------------------- geometry fingerprints (find identical STLs)
geo = {}
tri_cache = {}
for a in ALL:
    if a["ext"] == ".stl":
        tris = read_stl(a["path"])
        tri_cache[a["rel"]] = tris
        geo[a["rel"]] = geometry_key(tris)

# ---------------------------------------------------------------- per-row derived data
BG = re.compile(r"_0\.4n_(\d\.\d+)mm_PLA_MK4S_((?:\d+h)?\d+m)\.bgcode$")


def part_name(row, base):
    """'HitTargetTopV4-AdapterHolder.stl' -> 'AdapterHolder'."""
    for stem in row["stems"]:
        m = re.match(re.escape(stem) + r"[-_.]?(.*)", base)
        if m:
            rest = os.path.splitext(m.group(1))[0]
            rest = re.sub(r"_0\.4n_.*$", "", rest)
            return rest or base
    return base


for row in ROWS:
    files = row["files"]
    row["last_edit"] = max(a["mtime"] for a in files)
    hints = []
    bg = [a for a in files if a["ext"] == ".bgcode"]
    if bg:
        parts = []
        for a in sorted(bg, key=lambda x: x["base"]):
            m = BG.search(a["base"])
            label = part_name(row, a["base"]) if a["base"] else a["base"]
            parts.append(f"{label} {m.group(2) if m else ''}".strip())
        hints.append("Sliced for printing: " + "; ".join(parts))
    if any(a["sub"] == "MondayPrint" for a in files):
        hints.append("In your 'MondayPrint' print batch")
    if any(a["sub"] == "LaserCutterPinball" for a in files):
        hints.append("In your 'LaserCutterPinball' folder (Illustrator files for the laser cutter)")
    if any(a["sub"] == "Loading" for a in files):
        hints.append("Also in your 'Prototyping Pinball/Loading' folder")
    if any(a["sub"] == "" and a["rel"].startswith(PP) for a in files):
        hints.append("Also copied into your 'Prototyping Pinball' folder")
    if any(a["ext"] == ".ai" and not a["rel"].startswith(PP) for a in files):
        hints.append("Has an Illustrator (.ai) file, i.e. prepared for the laser cutter")
    if any(a["ext"] == ".3mf" for a in files):
        hints.append("Has a PrusaSlicer 3MF project")
    stl_files = [a for a in files if a["ext"] == ".stl"]
    empty_stl = [a for a in stl_files if geo.get(a["rel"]) == "empty"]
    if empty_stl:
        hints.append("Empty STL (no geometry): " + ", ".join(a["base"] for a in empty_stl))
    has_source = any(a["ext"] == ".fcstd" for a in files)
    if stl_files and not has_source:
        hints.append("No FreeCAD source file in the folder")
    row["hints"] = hints

# newest edit per component (only when the component has several rows)
VERSION_LINES = {"Playfield plate", "Flipper lever", "Flipper drive adapter", "Hit target (top)", "Hit target (mid)",
                 "Drop target", "Loading hook", "Servo adapter"}
by_comp = defaultdict(list)
for row in ROWS:
    by_comp[row["comp"]].append(row)
for comp, rows in by_comp.items():
    if len(rows) > 1 and comp in VERSION_LINES:
        newest = max(r["last_edit"] for r in rows)
        winners = [r for r in rows if r["last_edit"] == newest]
        if len(winners) == 1:
            winners[0]["hints"].append("Most recently edited version of this component")

# ---------------------------------------------------------------- previews
SIZE = 256


def fit_square(im, size=SIZE, bg=(246, 247, 249)):
    im = im.convert("RGB")
    im.thumbnail((size, size), Image.LANCZOS)
    canvas = Image.new("RGB", (size, size), bg)
    canvas.paste(im, ((size - im.width) // 2, (size - im.height) // 2))
    return canvas


def fcstd_thumb(row):
    cands = [a for a in row["files"] if a["ext"] == ".fcstd"]
    cands.sort(key=lambda a: (a["rel"].startswith(PP), a["rel"]))
    for a in cands:
        try:
            data = zipfile.ZipFile(a["path"]).read("thumbnails/Thumbnail.png")
            im = Image.open(io.BytesIO(data)).convert("RGBA")
            bg = Image.new("RGBA", im.size, (235, 235, 235, 255))
            bg.alpha_composite(im)
            return fit_square(bg)
        except Exception:
            continue
    return None


def part_preview(row):
    files = row["files"]
    stls = [a for a in files if a["ext"] == ".stl" and geo.get(a["rel"]) not in (None, "empty")]
    if stls:
        stls.sort(key=lambda a: -a["size"])
        # skip exact duplicates within the row
        seen, uniq = set(), []
        for a in stls:
            if geo[a["rel"]] not in seen:
                seen.add(geo[a["rel"]])
                uniq.append(a)
        row["preview_parts"] = [part_name(row, a["base"]) for a in uniq[:4]]
        row["preview_extra"] = max(0, len(uniq) - 4)
        tiles = [render_stl(tri_cache[a["rel"]], SIZE) for a in uniq[:4]]
        return montage(tiles, SIZE)
    pngs = [a for a in files if a["ext"] == ".png"]
    if pngs:
        return fit_square(Image.open(pngs[0]["path"]), bg=(255, 255, 255))
    svgs = sorted((a for a in files if a["ext"] == ".svg"), key=lambda a: -a["size"])
    valid = [(a, render_svg(a["path"], SIZE)) for a in svgs]
    valid = [(a, t) for a, t in valid if t is not None]
    if valid:
        top = valid[0][0]["size"]
        chosen = [(a, t) for a, t in valid if top < 5 * a["size"]][:4]
        row["preview_parts"] = [part_name(row, a["base"]) for a, _ in chosen]
        tiles = [t for _, t in chosen]
        return montage(tiles, SIZE) if len(tiles) > 1 else tiles[0]
    return None


for row in ROWS:
    row["thumb"] = fcstd_thumb(row)
    row["preview_parts"] = row.get("preview_parts", [])
    row["preview"] = part_preview(row)

print("rows with FreeCAD thumbnail:", sum(1 for r in ROWS if r["thumb"]), "| with part preview:", sum(1 for r in ROWS if r["preview"]),
      "| with neither:", [r["ver"] for r in ROWS if not r["thumb"] and not r["preview"]])

# ---------------------------------------------------------------- display strings for files
def display_file(a):
    name = a["base"]
    m = BG.search(name)
    if m:
        name = name[: m.start()] + f".bgcode  [{m.group(1)} mm, {m.group(2)}]"
    return (a["sub"] + "/ " if a["sub"] else (PP + "/ " if a["rel"].startswith(PP) else "")) + name


for row in ROWS:
    disp = [display_file(a) for a in sorted(row["files"], key=lambda x: (x["rel"].startswith(PP), x["ext"] != ".fcstd", x["base"]))]
    if len(disp) > 8:
        disp = disp[:7] + [f"+ {len(disp) - 7} more files"]
    row["files_text"] = "\n".join(disp)

# ---------------------------------------------------------------- workbook
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_small = Font(name=FONT, size=8, color="404040")
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color="1F2A44")
f_input = Font(name=FONT, size=10, bold=True, color="0000FF")
f_note = Font(name=FONT, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", start_color="1F2A44")
fill_input = PatternFill("solid", start_color="FFF2A8")
fill_band = PatternFill("solid", start_color="EEF2F8")
fill_used = PatternFill("solid", start_color="C6EFCE")
fill_not = PatternFill("solid", start_color="E3E3E3")
thin = Side(style="thin", color="C8CDD6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
wrap_mid = Alignment(wrap_text=True, vertical="center")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "Parts"

HEADERS = ["Component", "Version", "USED IN FINAL MACHINE?", "FreeCAD preview", "Part / outline preview", "What it is (my guess)",
           "Hints from your files", "Last edited", "Your notes", "Made how", "Group", "Files found"]
WIDTHS = [26, 24, 15, 18.5, 18.5, 40, 44, 12, 34, 13, 24, 54]
HEAD_ROW, FIRST = 4, 5
LAST = FIRST + len(ROWS) - 1
ROW_PT = 98  # 130 px, matches a 18.5-wide column (~135 px)

ws["A1"] = "Pinball machine - parts inventory (every version I found in your fablab folder)"
ws["A1"].font = f_title
ws["A2"] = "How: click a yellow cell in column C and pick  Used  or  Not used.  Blank = not decided yet.  Rows turn green (used) or grey (not used); the Summary tab counts for you."
ws["A2"].font = f_norm
ws["A3"] = "Tip: filter column C by 'Used' to get your final parts list.  Don't sort - the preview pictures stay where they are.  Hints are clues, not proof (a part can be sliced and still never used)."
ws["A3"].font = f_note

for c, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
    cell = ws.cell(row=HEAD_ROW, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = f_head, fill_head, center, border
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[HEAD_ROW].height = 34

comp_order, band = [], {}
for row in ROWS:
    if row["comp"] not in band:
        band[row["comp"]] = len(band) % 2 == 1
        comp_order.append(row["comp"])

for i, row in enumerate(ROWS):
    r = FIRST + i
    ws.row_dimensions[r].height = ROW_PT
    vals = [
        row["comp"], row["ver"], None, None, None, row["desc"],
        ("\n".join("- " + h for h in row["hints"]) if row["hints"] else "-"),
        row["last_edit"], None, row["how"], row["group"], row["files_text"],
    ]
    if row["preview_parts"] and len(row["preview_parts"]) > 1:
        vals[5] += "\nPreview tiles (top-left to bottom-right): " + ", ".join(row["preview_parts"]) + (f" (+{row['preview_extra']} more part{'s' if row['preview_extra'] > 1 else ''})" if row.get("preview_extra") else "")
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = f_norm
        cell.border = border
        cell.alignment = wrap_top if c in (6, 7, 9, 12) else wrap_mid
        if band[row["comp"]]:
            cell.fill = fill_band
    ws.cell(row=r, column=1).font = f_bold
    ws.cell(row=r, column=8).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=8).alignment = center
    ws.cell(row=r, column=12).font = f_small
    ws.cell(row=r, column=7).font = Font(name=FONT, size=9)
    ic = ws.cell(row=r, column=3)
    ic.fill, ic.font, ic.alignment = fill_input, f_input, center


def place(img_pil, col0, row0):
    bio = io.BytesIO()
    img_pil.save(bio, "PNG", optimize=True)
    bio.seek(0)
    xl = XLImage(bio)
    pad = 28575  # 3 px in EMU
    xl.anchor = TwoCellAnchor(
        editAs="twoCell",
        _from=AnchorMarker(col=col0, colOff=pad, row=row0, rowOff=pad),
        to=AnchorMarker(col=col0 + 1, colOff=-0 if False else 0, row=row0 + 1, rowOff=0),
    )
    ws.add_image(xl)


for i, row in enumerate(ROWS):
    if row["thumb"] is not None:
        place(row["thumb"], 3, FIRST + i - 1)
    if row["preview"] is not None:
        place(row["preview"], 4, FIRST + i - 1)

# dropdown, filters, freeze, conditional formats
dv = DataValidation(type="list", formula1='"Used,Not used"', allow_blank=True, showDropDown=False)
dv.promptTitle, dv.prompt = "Used in the final machine?", "Pick Used or Not used. Leave blank if you haven't decided."
dv.errorTitle, dv.error = "Pick from the list", "Choose Used or Not used (or clear the cell)."
dv.showInputMessage = dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add(f"C{FIRST}:C{LAST}")
ws.auto_filter.ref = f"A{HEAD_ROW}:{get_column_letter(len(HEADERS))}{LAST}"
ws.freeze_panes = f"D{FIRST}"
rng = f"A{FIRST}:{get_column_letter(len(HEADERS))}{LAST}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$C{FIRST}="Used"'], fill=fill_used, font=Font(name=FONT, color="0B5D1E")))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$C{FIRST}="Not used"'], fill=fill_not, font=Font(name=FONT, color="7A7A7A")))
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------- Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Pinball parts - summary (updates by itself as you click on the Parts tab)"
sm["A1"].font = f_title
rngA = f"Parts!$A${FIRST}:$A${LAST}"
rngC = f"Parts!$C${FIRST}:$C${LAST}"
kp = [
    ("Total part versions", f"=COUNTA({rngA})"),
    ("Marked Used", f'=COUNTIF({rngC},"Used")'),
    ("Marked Not used", f'=COUNTIF({rngC},"Not used")'),
    ("Not decided yet", "=B3-B4-B5"),
    ("Progress (decided / total)", "=IF(B3=0,0,(B4+B5)/B3)"),
]
for k, (lab, fml) in enumerate(kp, start=3):
    sm.cell(row=k, column=1, value=lab).font = f_bold
    c = sm.cell(row=k, column=2, value=fml)
    c.font = f_norm
    c.alignment = Alignment(horizontal="right")
    if lab.startswith("Progress"):
        c.number_format = "0%"
head2 = ["Group", "Component", "Versions", "Used", "Not used", "Not decided"]
HR2 = 10
for c, h in enumerate(head2, start=1):
    cell = sm.cell(row=HR2, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = f_head, fill_head, center, border
grp_of = {}
for row in ROWS:
    grp_of.setdefault(row["comp"], row["group"])
for k, comp in enumerate(comp_order):
    r = HR2 + 1 + k
    sm.cell(row=r, column=1, value=grp_of[comp])
    sm.cell(row=r, column=2, value=comp)
    sm.cell(row=r, column=3, value=f"=COUNTIF({rngA},B{r})")
    sm.cell(row=r, column=4, value=f'=COUNTIFS({rngA},B{r},{rngC},"Used")')
    sm.cell(row=r, column=5, value=f'=COUNTIFS({rngA},B{r},{rngC},"Not used")')
    sm.cell(row=r, column=6, value=f"=C{r}-D{r}-E{r}")
    for c in range(1, 7):
        cell = sm.cell(row=r, column=c)
        cell.font, cell.border = f_norm, border
        cell.alignment = Alignment(horizontal="center" if c > 2 else "left", vertical="center")
last2 = HR2 + len(comp_order)
tr = last2 + 1
sm.cell(row=tr, column=2, value="Total").font = f_bold
for c, col in zip(range(3, 7), "CDEF"):
    cell = sm.cell(row=tr, column=c, value=f"=SUM({col}{HR2 + 1}:{col}{last2})")
    cell.font, cell.alignment = f_bold, Alignment(horizontal="center")
for c in range(1, 7):
    sm.cell(row=tr, column=c).border = Border(top=Side(style="medium", color="1F2A44"))
sm.conditional_formatting.add(f"F{HR2 + 1}:F{last2}", FormulaRule(formula=[f"F{HR2 + 1}>0"], fill=PatternFill("solid", start_color="FFE9B3")))
for c, w in zip("ABCDEF", (30, 30, 11, 11, 11, 13)):
    sm.column_dimensions[c].width = w
sm.freeze_panes = f"A{HR2 + 1}"

# ---------------- Not pinball
npb = wb.create_sheet("Not pinball")
npb["A1"] = "Files I left out of the pinball list - tell me if I got one wrong"
npb["A1"].font = f_title
for c, h in enumerate(["File / folder", "What it is", "Why I left it out", "It IS pinball?"], start=1):
    cell = npb.cell(row=3, column=c, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = f_head, fill_head, center, border
for k, (a, b, c_) in enumerate(NON_PINBALL, start=4):
    for c, v in enumerate([a, b, c_, None], start=1):
        cell = npb.cell(row=k, column=c, value=v)
        cell.font, cell.border, cell.alignment = f_norm, border, wrap_mid
    npb.cell(row=k, column=4).fill = fill_input
    npb.cell(row=k, column=4).font = f_input
    npb.cell(row=k, column=4).alignment = center
    npb.row_dimensions[k].height = 30
dv2 = DataValidation(type="list", formula1='"Yes - it is pinball"', allow_blank=True)
npb.add_data_validation(dv2)
dv2.add(f"D4:D{3 + len(NON_PINBALL)}")
for c, w in zip("ABCD", (44, 50, 60, 20)):
    npb.column_dimensions[c].width = w

# ---------------- How to use
hw = wb.create_sheet("How to use")
hw["A1"] = "How to use this workbook"
hw["A1"].font = f_title
lines = [
    ("What to edit", "Only the yellow cells: column C on the Parts tab (Used / Not used), your notes in column I, and column D on the Not pinball tab. Everything else is filled in for you."),
    ("Example", "Component 'Hit target (top)': click 'Used' on V4 and 'Not used' on V1, V2, V3 and V5. The row turns green or grey and the Summary tab updates."),
    ("Blank cells", "A blank means 'not decided yet'. The Summary tab shows how many are still open."),
    ("Filter, don't sort", "Use the arrows in the header row (for example column C = Used). Sorting would separate the rows from their pictures."),
    ("Previews", "Left: the picture FreeCAD saved in the file (sometimes cropped or empty). Right: my render of the exported STL parts (up to 4 tiles, largest part first; names are listed in the description) or of the laser-cut SVG outline."),
    ("Hints column", "'Sliced for printing' = a .bgcode file exists (a print file for the Prusa MK4S). 'MondayPrint' = the file sits in your Prototyping Pinball/MondayPrint folder. 'Also copied into Prototyping Pinball' = you copied it into that folder, which looks like a final set. 'Most recently edited' = the newest save among that component's versions. These are clues, not proof."),
    ("What is my guess", "The descriptions come from file names, the previews, your Prototyping report and your LoadingAll.png diagram. Slovak words I decoded: Rucicka = flipper lever, Drevo = wood, Zabrana = barrier, Obmedzovac = limiter, Podlozka = washer, Trojuholnik = triangle, Uholnik = angle piece. Correct anything wrong in the notes column."),
    ("Scope", "Built from C:\\Users\\jakub\\Desktop\\KdG\\fablab (root files and the Prototyping Pinball folder). FreeCAD autosave backups (.FCBak) are ignored."),
]
for k, (a, b) in enumerate(lines, start=3):
    hw.cell(row=k, column=1, value=a).font = f_bold
    c = hw.cell(row=k, column=2, value=b)
    c.font, c.alignment = f_norm, wrap_top
    hw.cell(row=k, column=1).alignment = wrap_top
    hw.row_dimensions[k].height = 16 * max(1, -(-len(b) // 95))
hw.column_dimensions["A"].width = 22
hw.column_dimensions["B"].width = 100

wb.active = 0
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")

# expose for the verification step
import json  # noqa: E402

with open(os.path.join(os.path.dirname(__file__), "rows_summary.json"), "w", encoding="utf8") as fh:
    json.dump([dict(comp=r["comp"], ver=r["ver"], group=r["group"], hints=r["hints"], n_files=len(r["files"]),
                    thumb=r["thumb"] is not None, preview=r["preview"] is not None, last=str(r["last_edit"])) for r in ROWS], fh, indent=1)

# debug: dump previews so they can be inspected
_pd = os.path.join(os.path.dirname(__file__), "previews")
os.makedirs(_pd, exist_ok=True)
for _i, _r in enumerate(ROWS):
    for _k in ("thumb", "preview"):
        if _r[_k] is not None:
            _r[_k].save(os.path.join(_pd, f"{_i:02d}_{_k}.png"))
